"""
Gera o documento de entrega da extensão (.docx) com formatação adequada
para impressão e preenchimento manual dos campos.

Uso:
  py -3.12 scripts/gerar_doc_extensao.py
"""
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "Projeto_Extensao_EduMap_v3.docx"

UNDERLINE = "_______________________________________"
LONG_UNDER = "_______________________________________________________________"


def _shade_cell(cell, color_hex: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color_hex)
    tc_pr.append(shd)


def _set_cell_borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:color"), "888888")
        tc_borders.append(b)
    tc_pr.append(tc_borders)


def add_title(doc, text, level=0, color=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = True
    if level == 0:
        run.font.size = Pt(20)
        if color:
            run.font.color.rgb = color
    elif level == 1:
        run.font.size = Pt(15)
    elif level == 2:
        run.font.size = Pt(13)
    else:
        run.font.size = Pt(11)
    return p


def add_paragraph(doc, text, bold=False, italic=False, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    return p


def add_field_table(doc, rows):
    """Tabela 2 colunas: Campo | Informação. rows: list of (label, value)."""
    table = doc.add_table(rows=len(rows), cols=2)
    table.autofit = False
    for i, (label, value) in enumerate(rows):
        c0 = table.cell(i, 0)
        c1 = table.cell(i, 1)
        c0.width = Cm(6)
        c1.width = Cm(11)
        # Label
        c0.text = ""
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(label)
        r0.bold = True
        r0.font.size = Pt(10)
        _shade_cell(c0, "F3F4F6")
        # Value
        c1.text = ""
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(value)
        r1.font.size = Pt(10)
        _set_cell_borders(c0)
        _set_cell_borders(c1)
    doc.add_paragraph()


def add_signature_block(doc, lines):
    """Bloco de assinatura com underline visual."""
    for line in lines:
        p = doc.add_paragraph()
        run = p.add_run(line)
        run.font.size = Pt(10)
    doc.add_paragraph()


def add_bullet_list(doc, items, size=11):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.runs[0] if p.runs else p.add_run("")
        # Re-add run with text
        p.text = ""
        r = p.add_run(item)
        r.font.size = Pt(size)


def main():
    doc = Document()

    # Margens
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2)

    # Estilo padrão
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # ── CAPA ──────────────────────────────────────────────────────────────────
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(
        "Implementação de Tecnologia Assistiva\n"
        "para Diagnóstico de Aprendizagem e Capacitação Docente"
    )
    title_run.bold = True
    title_run.font.size = Pt(18)
    title_run.font.color.rgb = RGBColor(29, 78, 216)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub.add_run("Projeto de Extensão Universitária — EduMap")
    sub_run.italic = True
    sub_run.font.size = Pt(13)
    sub_run.font.color.rgb = RGBColor(107, 114, 128)

    doc.add_paragraph()

    # Identificação na capa
    add_title(doc, "Identificação", level=2)
    add_field_table(doc, [
        ("Aluno extensionista", "Willian Vidal Lima"),
        ("Curso", UNDERLINE),
        ("Instituição", UNDERLINE),
        ("Disciplina/Componente", "Atividade de Extensão Universitária"),
        ("Período letivo", UNDERLINE),
        ("Data da entrega", UNDERLINE),
    ])

    doc.add_page_break()

    # ── 1. RESUMO EXECUTIVO ────────────────────────────────────────────────────
    add_title(doc, "1. Resumo Executivo", level=1)
    add_paragraph(doc,
        "O EduMap é uma plataforma web gratuita que automatiza o diagnóstico "
        "pedagógico a partir de provas reais aplicadas a alunos. Combinando "
        "OCR (Reconhecimento Óptico de Caracteres) e uma taxonomia educacional "
        "hierárquica, o sistema lê a prova, classifica cada questão por área "
        "de conhecimento, nível cognitivo (Taxonomia de Bloom) e tópico "
        "específico, e gera relatórios que mostram exatamente onde cada aluno "
        "tem dificuldade — não apenas \"errou em Geometria\", mas \"errou na "
        "classificação de triângulos pelos lados\".")
    add_paragraph(doc,
        "A ferramenta foi desenvolvida no contexto desta extensão universitária "
        "e prestada como serviço gratuito a professores particulares, centros "
        "de reforço e escolas de bairro, otimizando o tempo de ensino e "
        "direcionando o esforço pedagógico ao que importa.")

    # ── 2. APRESENTAÇÃO ────────────────────────────────────────────────────────
    add_title(doc, "2. Apresentação do Projeto", level=1)
    add_field_table(doc, [
        ("Nome", "EduMap"),
        ("Tipo", "Plataforma web (frontend + backend) com IA assistiva"),
        ("URL pública", "https://edumap-frontend-nu.vercel.app"),
        ("Repositório (frontend)", "https://github.com/Nerodowble/edumap_frontend"),
        ("Repositório (backend)", "https://github.com/Nerodowble/edumap_ia"),
        ("Custo para o usuário", "R$ 0,00 (totalmente gratuito)"),
        ("Acesso", "cadastro próprio, sem instalação local"),
    ])

    # ── 3. JUSTIFICATIVA ───────────────────────────────────────────────────────
    add_title(doc, "3. Justificativa", level=1)
    add_title(doc, "O problema", level=3)
    add_paragraph(doc,
        "Professores da educação básica e do reforço escolar enfrentam uma "
        "limitação prática diária: ao corrigir uma prova, identificam que o "
        "aluno errou, mas raramente conseguem mapear sistematicamente onde "
        "está a lacuna — qual habilidade específica precisa ser reforçada.")
    add_paragraph(doc, "Isso leva a:", bold=True)
    add_bullet_list(doc, [
        "Aulas de revisão pouco direcionadas",
        "Reforço genérico que repete o que o aluno já sabe",
        "Tempo perdido com avaliação manual de padrões de erro",
        "Dificuldade de individualizar o cuidado pedagógico",
    ])

    add_title(doc, "O que o EduMap resolve", level=3)
    add_bullet_list(doc, [
        "Automatiza a leitura das provas (OCR), eliminando trabalho manual",
        "Classifica taxonomicamente cada questão até o conceito específico",
        "Gera relatórios apontando os pontos críticos da turma e de cada aluno",
        "Sugere ações pedagógicas baseadas nos dados",
    ])

    # ── 4. OBJETIVOS ──────────────────────────────────────────────────────────
    add_title(doc, "4. Objetivos", level=1)
    add_title(doc, "Geral", level=3)
    add_paragraph(doc,
        "Disponibilizar gratuitamente uma ferramenta de diagnóstico pedagógico "
        "automatizado a professores que atuam fora dos sistemas educacionais "
        "com recursos para tecnologia, ampliando a precisão do cuidado "
        "individualizado oferecido aos alunos.")
    add_title(doc, "Específicos", level=3)
    add_bullet_list(doc, [
        "Construir uma plataforma web acessível (qualquer navegador, sem instalação)",
        "Implementar OCR de provas em PDF e imagem",
        "Estruturar uma taxonomia educacional hierárquica (até 6 níveis)",
        "Permitir que professores cadastrem turmas, alunos e respostas",
        "Produzir relatórios em formato visual e em PDF para compartilhamento",
        "Manter o código aberto, permitindo reuso e adaptação",
    ])

    # ── 5. METODOLOGIA ─────────────────────────────────────────────────────────
    add_title(doc, "5. Metodologia e Funcionamento", level=1)
    add_title(doc, "Tecnologias utilizadas", level=3)
    add_bullet_list(doc, [
        "Backend: Python (FastAPI), Tesseract OCR, PyMuPDF, PostgreSQL",
        "Frontend: Next.js 14 (React, TypeScript), Tailwind CSS",
        "Infraestrutura: Render (backend + banco), Vercel (frontend)",
        "Autenticação: JWT com 3 níveis (admin geral, admin escolar, professor)",
    ])

    add_title(doc, "Fluxo de uso pelo professor", level=3)
    add_bullet_list(doc, [
        "Cadastro próprio na plataforma (gratuito)",
        "Criação de turma e cadastro de alunos",
        "Upload da prova (PDF/foto)",
        "Sistema executa OCR e classifica cada questão",
        "Professor define o gabarito oficial",
        "Professor lança as respostas dos alunos",
        "Sistema gera relatórios (visão geral, drilldown, por aluno, PDF completo)",
    ])

    add_title(doc, "Taxonomia educacional", level=3)
    add_paragraph(doc,
        "A plataforma vem com taxonomias prontas para múltiplas etapas:")
    add_bullet_list(doc, [
        "Ensino Fundamental II (282 nós): Matemática (com Geometria detalhada), "
        "Português, Geografia, Ciências, História, Inglês, Artes, Ed. Física",
        "Ensino Superior (361 nós): Psicologia, Saúde Mental, SUS",
    ])
    add_paragraph(doc,
        "A taxonomia é totalmente editável pelo administrador via interface "
        "gráfica, permitindo expansão para novas matérias, etapas e cursos.")

    # ── 6. IMPACTO SOCIAL ─────────────────────────────────────────────────────
    add_title(doc, "6. Impacto Social", level=1)
    add_title(doc, "Beneficiários diretos", level=3)
    add_bullet_list(doc, [
        "Professores particulares e de reforço escolar que muitas vezes "
        "trabalham sozinhos, sem equipe pedagógica de apoio",
        "Pequenas escolas de bairro com recursos tecnológicos limitados",
        "Educadores em formação (estagiários, monitores) que ganham uma "
        "ferramenta de análise rica para usar em sua prática",
    ])

    add_title(doc, "Beneficiários indiretos", level=3)
    add_bullet_list(doc, [
        "Alunos: recebem reforço mais direcionado e eficiente",
        "Famílias: podem apoiar em casa de forma específica",
        "Sistema educacional como um todo: cultura de diagnóstico baseado em dados",
    ])

    add_title(doc, "Características que reforçam o caráter de extensão", level=3)
    add_bullet_list(doc, [
        "Gratuito — sem assinatura, paywall ou limite de uso",
        "Aberto — código-fonte público, permite extensões em trabalhos futuros",
        "Acessível — funciona em qualquer navegador, inclusive em celulares",
        "Em português — interface inteiramente em PT-BR",
        "Sem coleta abusiva de dados — apenas o necessário para o funcionamento",
    ])

    doc.add_page_break()

    # ── 7. ATIVIDADE PRÁTICA ──────────────────────────────────────────────────
    add_title(doc, "7. Atividade Prática Realizada", level=1)
    add_title(doc, "Sessão de diagnóstico com professor parceiro", level=3)
    add_field_table(doc, [
        ("Data da ação", "____ / ____ / ________"),
        ("Horário (início e término)", "____:____ até ____:____"),
        ("Carga horária da sessão presencial", "8 hora(s)"),
        ("Local", LONG_UNDER),
        ("Endereço completo", LONG_UNDER),
        ("Geolocalização (latitude, longitude)", LONG_UNDER),
    ])

    add_title(doc, "Descrição da atividade", level=3)
    add_paragraph(doc,
        "Apresentação do sistema EduMap ao(à) professor(a) parceiro(a), "
        "demonstração do fluxo completo (cadastro → upload de prova → "
        "classificação automática → relatório), aplicação em uma prova real "
        "fornecida pelo(a) docente, análise conjunta dos resultados e coleta "
        "de feedback e validação pedagógica.")

    add_title(doc, "Carga horária total do projeto de extensão", level=3)
    add_paragraph(doc,
        "A carga horária total dedicada a este projeto de extensão soma "
        "200 horas (duzentas horas), distribuídas conforme abaixo:", bold=True)
    add_field_table(doc, [
        ("Pesquisa de campo e estudo prévio",
         "≈ 30h — levantamento de necessidades de professores, estudo da "
         "Taxonomia de Bloom e estruturação da taxonomia educacional"),
        ("Desenvolvimento do projeto (backend + frontend + IA)",
         "≈ 140h — programação da plataforma, OCR, classificadores, "
         "interface, banco de dados, autenticação, painel administrativo"),
        ("Testes unitários e de integração",
         "≈ 22h — testes automatizados (44 casos), testes manuais e "
         "validação do classificador taxonômico"),
        ("Sessão prática presencial com professor parceiro",
         "8h — demonstração, aplicação em prova real, análise conjunta e "
         "devolutiva pedagógica (data registrada acima)"),
        ("Total", "200 horas"),
    ])

    add_title(doc, "Critérios de avaliação atendidos", level=3)
    add_bullet_list(doc, [
        "[ X ] Atividade executada na prática com docente parceiro(a)",
        "[ X ] Registro fotográfico anexado (extensionista presente nas fotos)",
        "[ X ] Geolocalização registrada (campo acima)",
        "[ X ] Atestado de presencialidade assinado (próximas seções)",
        "[ X ] Revisão completa das informações antes da finalização",
    ])

    # ── 8. RESULTADOS ─────────────────────────────────────────────────────────
    add_title(doc, "8. Resultados Observados", level=1)
    add_paragraph(doc,
        "Durante a sessão, o(a) professor(a) parceiro(a) testou o sistema em "
        "uma prova real e validou:")
    add_bullet_list(doc, [
        "A correção do OCR na leitura das questões",
        "A precisão da classificação taxonômica das questões",
        "A clareza dos relatórios gerados (interface e PDF)",
        "A utilidade prática da ferramenta para o cotidiano docente",
        "As recomendações pedagógicas geradas automaticamente",
    ])
    add_paragraph(doc,
        "A devolutiva qualitativa do(a) docente está registrada na seção 11 "
        "deste documento.")

    # ── 9. CONSIDERAÇÕES FINAIS ───────────────────────────────────────────────
    add_title(doc, "9. Considerações Finais e Continuidade", level=1)
    add_paragraph(doc,
        "O EduMap nasceu como projeto de extensão universitária mas foi "
        "construído desde o início com arquitetura escalável e código aberto, "
        "permitindo:")
    add_bullet_list(doc, [
        "Uso continuado pelos professores parceiros após o término da extensão",
        "Estudantes futuros estenderem o projeto (novas taxonomias, matérias, funcionalidades)",
        "Possibilidade de virar produto/serviço institucional, mantendo a versão gratuita aberta",
    ])
    add_paragraph(doc,
        "A ferramenta permanecerá no ar e acessível em "
        "https://edumap-frontend-nu.vercel.app, e o código continuará mantido "
        "publicamente nos repositórios listados na seção 2.")

    doc.add_page_break()

    # ── 10. EXTENSIONISTA ─────────────────────────────────────────────────────
    add_title(doc, "10. Identificação do Aluno Extensionista", level=1)
    add_field_table(doc, [
        ("Nome completo", "Willian Vidal Lima"),
        ("CPF", UNDERLINE),
        ("RA / Matrícula", UNDERLINE),
        ("Curso", UNDERLINE),
        ("Instituição", UNDERLINE),
        ("E-mail", "willian.lima@legalbot.com.br"),
    ])

    doc.add_paragraph()
    add_title(doc, "ASSINATURA DO EXTENSIONISTA", level=2,
              color=RGBColor(29, 78, 216))

    sig_table_aluno = doc.add_table(rows=1, cols=1)
    sig_cell_aluno = sig_table_aluno.cell(0, 0)
    _set_cell_borders(sig_cell_aluno)
    _shade_cell(sig_cell_aluno, "EFF6FF")

    p_aluno1 = sig_cell_aluno.paragraphs[0]
    p_aluno1.add_run("\n\n").font.size = Pt(10)
    p_aluno1 = sig_cell_aluno.add_paragraph()
    p_aluno1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_a = p_aluno1.add_run("_" * 60)
    r_a.font.size = Pt(14)

    p_aluno2 = sig_cell_aluno.add_paragraph()
    p_aluno2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_b = p_aluno2.add_run("Willian Vidal Lima")
    r_b.bold = True
    r_b.font.size = Pt(11)

    p_aluno3 = sig_cell_aluno.add_paragraph()
    p_aluno3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_c = p_aluno3.add_run("Aluno extensionista")
    r_c.italic = True
    r_c.font.size = Pt(9)
    r_c.font.color.rgb = RGBColor(107, 114, 128)

    p_aluno4 = sig_cell_aluno.add_paragraph()
    p_aluno4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_d = p_aluno4.add_run("Data: ____ / ____ / ________")
    r_d.font.size = Pt(10)
    sig_cell_aluno.add_paragraph()

    doc.add_page_break()

    # ── 11. ATESTADO DO PROFESSOR PARCEIRO ────────────────────────────────────
    add_title(doc, "11. Atestado de Presencialidade — Professor(a) Parceiro(a)", level=1)

    decl = doc.add_paragraph()
    decl.paragraph_format.left_indent = Cm(0.5)
    decl.paragraph_format.right_indent = Cm(0.5)
    decl_run = decl.add_run(
        "Eu, abaixo identificado(a), atesto que o(a) aluno(a) Willian Vidal "
        "Lima realizou comigo, na data e local abaixo informados, uma sessão "
        "prática de demonstração e validação da plataforma EduMap, no "
        "contexto da sua atividade de extensão universitária. Aprovo o uso "
        "pedagógico da ferramenta apresentada e confirmo a presença física "
        "do extensionista durante toda a ação."
    )
    decl_run.italic = True
    decl_run.font.size = Pt(10)

    add_title(doc, "Dados do(a) Professor(a) Parceiro(a)", level=3)
    add_field_table(doc, [
        ("Nome completo", LONG_UNDER),
        ("CPF", UNDERLINE),
        ("Formação / Titulação", LONG_UNDER),
        ("Cargo / Função", LONG_UNDER),
        ("Instituição em que atua", LONG_UNDER),
        ("E-mail", LONG_UNDER),
        ("Telefone", UNDERLINE),
    ])

    add_title(doc, "Dados da Ação", level=3)
    add_field_table(doc, [
        ("Local da ação", LONG_UNDER),
        ("Endereço completo", LONG_UNDER),
        ("Geolocalização", LONG_UNDER),
        ("Data", "____ / ____ / ________"),
        ("Horário (início — término)", "____:____ até ____:____"),
        ("Carga horária da sessão presencial", "8 hora(s)"),
        ("Carga horária total do projeto", "200 horas (declaradas pelo extensionista)"),
    ])

    add_title(doc, "Descrição da Atividade", level=3)
    add_paragraph(doc,
        "Aplicação prática do sistema EduMap — plataforma web gratuita de "
        "diagnóstico pedagógico via OCR e análise por taxonomia de erros — em "
        "provas reais aplicadas pelo(a) professor(a). A sessão envolveu:")
    add_bullet_list(doc, [
        "Demonstração completa da plataforma",
        "Cadastro de turma e alunos pelo(a) próprio(a) docente",
        "Upload de prova(s) e leitura automática via OCR",
        "Classificação taxonômica das questões",
        "Lançamento de respostas e geração de relatórios",
        "Análise conjunta dos resultados e devolutiva do(a) professor(a)",
    ])

    add_title(doc, "Devolutiva e Aprovação Pedagógica", level=3)
    add_paragraph(doc,
        "Espaço para o(a) professor(a) parceiro(a) registrar comentários, "
        "sugestões e validação pedagógica do uso da ferramenta:", italic=True, size=10)

    for _ in range(8):
        add_paragraph(doc, LONG_UNDER + LONG_UNDER, size=10)

    # ── BLOCO DE ASSINATURA FORMAL (DESTACADO) ──
    doc.add_paragraph()
    add_title(doc, "ASSINATURA DO(A) PROFESSOR(A) PARCEIRO(A)", level=2,
              color=RGBColor(220, 38, 38))
    add_paragraph(doc,
        "Declaro estar ciente do conteúdo deste documento e atesto a veracidade "
        "das informações aqui prestadas. Confirmo a presença física do aluno "
        "extensionista Willian Vidal Lima durante toda a sessão prática descrita.",
        italic=True)

    # Caixa visual em volta da assinatura
    sig_table = doc.add_table(rows=1, cols=1)
    sig_cell = sig_table.cell(0, 0)
    _set_cell_borders(sig_cell)
    _shade_cell(sig_cell, "FFFBEB")
    sig_cell.text = ""

    # Linha grande para assinatura à mão
    p1 = sig_cell.paragraphs[0]
    p1.add_run("\n\n").font.size = Pt(10)
    p1 = sig_cell.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p1.add_run("_" * 60)
    r1.font.size = Pt(14)

    p2 = sig_cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("Assinatura do(a) Professor(a) Parceiro(a)")
    r2.bold = True
    r2.font.size = Pt(11)

    p3 = sig_cell.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = p3.add_run("(a mão, com nome completo legível,  OU  digitalmente via GOV.BR)")
    r3.italic = True
    r3.font.size = Pt(9)
    r3.font.color.rgb = RGBColor(107, 114, 128)

    sig_cell.add_paragraph()  # espaço

    p4 = sig_cell.add_paragraph()
    r4 = p4.add_run("Nome completo (em letra de forma): ")
    r4.bold = True
    r4.font.size = Pt(10)
    p4.add_run("_" * 50).font.size = Pt(10)

    p5 = sig_cell.add_paragraph()
    r5 = p5.add_run("CPF: ")
    r5.bold = True
    r5.font.size = Pt(10)
    p5.add_run("___________________________  ").font.size = Pt(10)
    r5b = p5.add_run("RG: ")
    r5b.bold = True
    r5b.font.size = Pt(10)
    p5.add_run("___________________________").font.size = Pt(10)

    p6 = sig_cell.add_paragraph()
    r6 = p6.add_run("Data da assinatura: ")
    r6.bold = True
    r6.font.size = Pt(10)
    p6.add_run("____ / ____ / ________").font.size = Pt(10)
    sig_cell.add_paragraph()

    doc.add_paragraph()
    obs = doc.add_paragraph()
    obs_run = obs.add_run(
        "⚠ Observação importante: a faculdade aceita assinatura manuscrita "
        "(com nome completo legível) OU assinatura digital via GOV.BR. Caso "
        "opte pelo GOV.BR, anexe o PDF assinado a esta entrega ou inclua a "
        "URL/código de validação no espaço acima."
    )
    obs_run.italic = True
    obs_run.font.size = Pt(9)
    obs_run.font.color.rgb = RGBColor(107, 114, 128)

    doc.add_page_break()

    # ── ANEXOS ────────────────────────────────────────────────────────────────
    add_title(doc, "Anexos", level=1)
    add_paragraph(doc,
        "A entrega completa deste projeto inclui, além deste documento:")
    add_bullet_list(doc, [
        "Fotografias da ação com o extensionista presente",
        "Captura de geolocalização (print do mapa ou coordenadas)",
        "Este documento preenchido e assinado",
        "(Opcional) Print do relatório gerado pela plataforma durante a sessão de teste",
    ])

    # Rodapé
    doc.add_paragraph()
    foot = doc.add_paragraph()
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = foot.add_run(
        "EduMap — projeto de extensão universitária · diagnóstico educacional · código aberto"
    )
    fr.italic = True
    fr.font.size = Pt(9)
    fr.font.color.rgb = RGBColor(107, 114, 128)

    doc.save(str(OUT))
    print(f"[OK] Documento salvo em: {OUT}")
    print(f"     Tamanho: {OUT.stat().st_size / 1024:.1f} KB")


if __name__ == "__main__":
    main()
